import sys; from xl import load
import openpyxl,warnings; warnings.filterwarnings('ignore')
f=sys.argv[1]; d=load(f)
wb=openpyxl.load_workbook(f,data_only=True)
for s,t in d.items():
    ws=wb[s if s in wb.sheetnames else s+' ']
    hdr=[h for h in t['hdr']]
    home=[h for h in hdr if 'Amount' in h]
    home=home[0] if home else None
    # rates
    if s.startswith('Data'): zr=ws['E2'].value; fx=None
    elif home: fx=ws['D2'].value; zr=ws['F1'].value
    else: zr=ws['E1'].value; fx=None
    n=0; flags=[]; hard=[]
    for r in t['rows']:
        v={k.strip():x for k,x in r['vals'].items()}; fm={k.strip():x for k,x in r['forms'].items()}
        if v['Unique ID'] is None: continue
        n+=1
        u=v['Total US$']; z=v['Total ZAR']; h=v.get(home.strip()) if home else None
        for col,val in (('Total US$',fm['Total US$']),('Total ZAR',fm['Total ZAR'])):
            if not (isinstance(val,str) and val.startswith('=')): hard.append((v['Unique ID'],col,val))
        if isinstance(u,(int,float)) and isinstance(z,(int,float)):
            if abs(z-u*zr)>1: flags.append((v['Unique ID'],'ZAR!=USD*rate',u,z,round(z/u,4) if u else None))
            if abs(z-u)<1: flags.append((v['Unique ID'],'ZAR==USD',u,z))
        else:
            flags.append((v['Unique ID'],'non-numeric',u,z))
        if home and isinstance(h,(int,float)) and isinstance(u,(int,float)):
            if abs(u-h*fx)>1: flags.append((v['Unique ID'],f'USD!=home*{fx}',h,u,round(u/h,4) if h else None))
            if abs(u-h)<1: flags.append((v['Unique ID'],'USD==home',h,u))
            if isinstance(z,(int,float)) and abs(z-h)<1: flags.append((v['Unique ID'],'ZAR==home',h,z))
        elif home: flags.append((v['Unique ID'],'home non-numeric',h,u))
    print(f'{s}: rows {n} zar_rate {zr} fx {fx} home {home}; hardcoded amount cells {len(hard)}')
    for x in hard: print('   hard',x)
    for x in flags: print('   FLAG',x)
