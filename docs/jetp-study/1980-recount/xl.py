import openpyxl, warnings, json, sys
from openpyxl.utils import range_boundaries
warnings.filterwarnings('ignore')
SKIP={'Dropdowns','PivotTables','Analysis - Dashboard','Financing Status'}
def load(f):
    wv=openpyxl.load_workbook(f,data_only=True); wf=openpyxl.load_workbook(f,data_only=False)
    out={}
    for ws in wv.worksheets:
        if ws.title in SKIP: continue
        wsf=wf[ws.title]
        t=list(wsf.tables.values())[0]
        c1,r1,c2,r2=range_boundaries(t.ref)
        hdr=[str(ws.cell(r1,c).value).strip() for c in range(c1,c2+1)]
        rows=[]
        for r in range(r1+1,r2+1):
            vals={h:ws.cell(r,c).value for h,c in zip(hdr,range(c1,c2+1))}
            forms={h:wsf.cell(r,c).value for h,c in zip(hdr,range(c1,c2+1))}
            if all(v is None for v in vals.values()): continue
            rows.append(dict(row=r,vals=vals,forms=forms,fill={h:ws.cell(r,c).fill.fgColor.rgb for h,c in zip(hdr,range(c1,c2+1))}))
        out[ws.title.strip()]=dict(hdr=hdr,rows=rows,totals_row=t.totalsRowCount)
    return out
